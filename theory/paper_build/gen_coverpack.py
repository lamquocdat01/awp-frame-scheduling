# -*- coding: utf-8 -*-
"""
gen_coverpack.py — Bộ submission ĐỒNG BỘ v2 (single author, biên đúng r=β(1−c)/(1−β)):
  03_Cover_Letter_v2.docx, 03_Highlights.docx, 03_Response_Template.docx, 00_Submission_Checklist.docx
Xuất vào submission_package/.
"""
from docx import Document
from docx.shared import Pt

PKG = ("/sessions/bold-wizardly-cannon/mnt/PhD Program/02. Working/"
       "03_AwP_Frame_Scheduling/submission_package/")
TITLE = ("Learning-Augmented Frame Scheduling: Consistency–Robustness Bounds for Gating "
         "with Untrusted Motion Predictors")


def newdoc():
    d = Document(); s = d.styles["Normal"]; s.font.name = "Times New Roman"; s.font.size = Pt(11)
    return d
def P(d, t, b=False, i=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; r.italic = i; return p


# ---------------- Cover letter ----------------
d = newdoc()
P(d, "Lam Quoc Dat", b=True)
P(d, "CERTES dissertation · [Affiliation]")
P(d, "[e-mail withheld in repository]")
P(d, "June 12, 2026")
P(d, "")
P(d, "To the Editor-in-Chief, IEEE Transactions on Computers", b=True)
P(d, "")
P(d, "Dear Editor,")
P(d, "I submit for your consideration the manuscript, “" + TITLE + "”, for publication in IEEE "
      "Transactions on Computers. I am the sole author.")
P(d, "Energy-constrained edge devices must decide online which video frames to spend a detector on. "
      "The prevailing wisdom is that content-aware gating — using a cheap motion cue to pick "
      "event-bearing frames — beats blind periodic sampling. A recent multi-protocol evaluation "
      "found the opposite: at equal energy, learned and hand-crafted gates are dominated by an "
      "oblivious periodic schedule, and the within-dataset ranking reverses cross-dataset. This "
      "manuscript explains the phenomenon at the theorem level and turns an empirical negative "
      "result into a structural impossibility theorem.")
P(d, "Specific contributions:", b=True)
for t in [
    "A formal model of gated perception sampling (windowed detection-recall, hard activation budget) "
    "in the algorithms-with-predictions framework, with a deployable prediction-error metric η on "
    "motion cues.",
    "A regime-indexed lower bound: a coverage bound for dense events and a new advice-commitment "
    "bound (Theorem 1′) for the high-budget regime a ≥ 1/(W+1), proving the advice arm contributes "
    "zero worst-case recall — the first impossibility result for content-aware sampling.",
    "A TIGHT consistency–robustness frontier for the deployed class of non-adaptive gates "
    "(Theorem 2): r = β(1−c)/(1−β), with HEDGE-GATE shown Pareto-optimal in that class (Theorem 3); "
    "equivalently robustness ratio Ω(1/ε) for (1+ε)-consistency.",
    "HEDGE-GATE, a budget-splitting meta-policy that attains the frontier and degrades affinely in η "
    "(Proposition 4), giving operators a certifiable trust dial; a feedback-aware variant is "
    "identified as a positive future direction.",
    "Empirical grounding: measured P-A anchors, a calibrated synthetic sweep validating the bounds, "
    "and a real-data run over a 121-video corpus (BMC, LASIESTA, CDnet2014) in which a motion cue of "
    "onset-AUC ≈ 0.51 yields a gate trailing the periodic schedule by 0.42–0.44 recall at low "
    "activation, with non-overlapping confidence intervals.",
]:
    d.add_paragraph(t, style="List Bullet")
P(d, "The work is original, not under review elsewhere, and I declare no conflicts of interest. It "
      "complements a companion study on certifiable rate scheduling. I believe it fits the scope of "
      "IEEE Transactions on Computers (energy-efficient edge computation, online scheduling with "
      "performance guarantees) and will interest both the systems and the online-algorithms "
      "communities.")
P(d, "Thank you for your consideration.")
P(d, ""); P(d, "Sincerely,"); P(d, "Lam Quoc Dat")
d.save(PKG + "03_Cover_Letter_v2.docx")

# ---------------- Highlights ----------------
d = newdoc()
d.add_heading("Highlights — " + TITLE, level=1)
for t in [
    "Content-aware frame gating is recast as an online algorithm consulting an untrusted motion "
    "predictor; its empirical collapse is the robustness arm of a provable trade-off.",
    "New advice-commitment lower bound (Theorem 1′) is non-vacuous in the high-budget regime where "
    "the standard coverage bound fails: the advice arm yields zero worst-case recall.",
    "Tight consistency–robustness frontier r = β(1−c)/(1−β) for the deployed class of non-adaptive "
    "gates; HEDGE-GATE is Pareto-optimal in that class.",
    "HEDGE-GATE attains the frontier and degrades affinely in prediction error η — a certifiable "
    "dial between exploiting predictions and guaranteeing recall.",
    "Validated on a 121-video real corpus (BMC, LASIESTA, CDnet2014): a motion cue (onset-AUC ≈ "
    "0.51) trails the periodic schedule by 0.42–0.44 recall at low activation.",
]:
    d.add_paragraph(t, style="List Bullet")
d.save(PKG + "03_Highlights.docx")

# ---------------- Response template ----------------
d = newdoc()
d.add_heading("Response to Reviewers — " + TITLE, level=1)
P(d, "I thank the reviewers for their time. Below I address each comment point by point; reviewer "
      "text is in italics, my response in regular type, and manuscript changes are noted with "
      "section/line references. [Fill after receiving reviews.]", i=True)
for rv in ["Reviewer 1", "Reviewer 2", "Reviewer 3"]:
    d.add_heading(rv, level=2)
    for k in ["1", "2", "3"]:
        P(d, "Comment %s.%s — [paste reviewer comment]" % (rv.split()[-1], k), i=True)
        P(d, "Response. [response; reference §__, p.__; change made: __]"); P(d, "")
d.add_heading("Anticipated question: scope of Theorems 2–3", level=2)
P(d, "Theorems 2–3 are proven for non-adaptive gates, which is exactly the deployed class "
      "(threshold S/RF gates and HEDGE-GATE). §V remarks that feedback-aware adaptive gates can "
      "exceed the frontier; this is stated as a positive design implication and future work, not a gap.")
d.save(PKG + "03_Response_Template.docx")

# ---------------- Submission checklist ----------------
d = newdoc()
d.add_heading("Submission Checklist — Topic 03", level=1)
P(d, "Target: IEEE Transactions on Computers. Sole author: Lam Quoc Dat ([e-mail withheld in repository]).", i=True)
d.add_heading("Package contents (submission_package/)", level=2)
for t in ["03_Paper_Draft_v2.docx / .pdf — manuscript (theory synced; proofs in Appendix A)",
          "03_Cover_Letter_v2.docx — cover letter (single author)",
          "03_Highlights.docx — 5 highlights",
          "03_Response_Template.docx — point-by-point response skeleton",
          "figures/ — fig1, fig3, fig4, fig5",
          "../03_Proofs_Theorems_Bilingual_EN_VN.docx/pdf — full bilingual proofs (supporting)",
          "../theory_supplement.md — synced theory note",
          "../results/ — CSVs, metrics_summary*.json, figures (reproducibility)"]:
    d.add_paragraph(t, style="List Bullet")
d.add_heading("Before you submit — author actions", level=2)
for t in ["VERIFY the proofs of Lemma 1, Theorem 1′, Theorem 2 and Proposition 4 (checklists in "
          "theory_supplement.md §5 and the bilingual proofs doc).",
          "Fill the [Affiliation] field in the title block and cover letter.",
          "Theorem 1 (dense regime) proof is now included (counting argument) — verify it with the others.",
          "Optional: run the uncapped 121-video real-data pass (MAX_FRAMES=None) to tighten §VI CIs.",
          "De-anonymize references [4],[5] per the venue policy, or keep anonymized if double-blind.",
          "Format pass to the IEEE TC template; figures ≥ 300 dpi (already); author/affiliation, "
          "funding and conflict statements."]:
    d.add_paragraph(t, style="List Bullet")
d.add_heading("Theory status (state honestly in the paper)", level=2)
for t in ["PROVEN: Lemma 1; Theorem 1′; Theorem 2 & 3 (non-adaptive class); Proposition 4 (under "
          "Assumption A).",
          "OPEN / future work: optimal feedback-aware (adaptive) gate."]:
    d.add_paragraph(t, style="List Bullet")
d.save(PKG + "00_Submission_Checklist.docx")
print("Saved cover letter v2, highlights, response template, checklist ->", PKG)
