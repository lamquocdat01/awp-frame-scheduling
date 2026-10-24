# -*- coding: utf-8 -*-
"""v3_package.py -- assemble submission_package/EJOR/ from Revision_v3 (run code/v3_compile.sh first).

Contents: 01_manuscript_paper_v3.pdf, 02_source_flat.zip (flat LaTeX source, verified to compile in an empty
folder), 03_S1_theory_supplement.pdf, 04_S2_empirical_supplement.pdf, 05_cover_letter.pdf, 06_highlights.docx,
07_abstract_keywords_ASCII.txt, MANIFEST.md (sha256). No file is edited by hand.
"""
import hashlib
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from v3_io import atomic_copy, atomic_open, atomic_write_text  # noqa: E402  (atomic writes, C7)

TOPIC = Path(__file__).resolve().parents[1]
REV = TOPIC / "Revision_v3"
PKG = TOPIC / "submission_package" / "EJOR"
REPO = Path(os.environ.get("AWP_REPO", r"D:\Repos\awp-frame-scheduling"))
TAG = "v3.0.1-submission"
# every derived data file the paper / supplements rely on (C7(a): hash now, re-hash the raw GitHub files after push)
DATA = ["Revision_v3/numbers.json", "Revision_v3/numbers.tex", "Revision_v3/S2_percategory_v3.csv",
        "results/v3/a8a_summary.json", "results/v3/a8a_ceiling.csv", "results/v3/a8a_sv_worstcase.csv",
        "results/v3/a8a_sv_family_all.csv", "results/v3/a8c_summary.json", "results/v3/a8c_alloc.csv",
        "results/v3/a8c_T2.csv", "results/v3/thm3_hedge_audit.txt", "results/v3/corruption_sweep_REAL_blockaware.csv",
        "results/v3/sweep_blockaware_summary.json", "results/v3/corruption_sweep_blockaware.csv",
        "results/v3/sweep_blockaware_synthetic_summary.json", "results/percategory.csv",
        "results/metrics_summary_REAL.json", "results/corruption_sweep_REAL.csv", "results/metrics_summary.json",
        "results/frontier_general.csv", "results/learned_cue_probe.json", "results/anchors_pa.csv",
        "energy/fit.json"]


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def flat_source(tmp: Path) -> list[Path]:
    files = []
    tex = (REV / "paper_v3.tex").read_text(encoding="utf-8")
    tex = tex.replace("{figures/", "{")
    # drop the template-provenance comments (internal paths) from the submitted source
    tex = "\n".join(l for l in tex.splitlines() if not l.startswith("%% ")) + "\n"
    atomic_write_text(tmp / "paper_v3.tex", tex, encoding="utf-8")
    files.append(tmp / "paper_v3.tex")
    for f in ("numbers.tex", "tab_a8a.tex", "tab_a8c.tex", "refs.bib", "paper_v3.bbl"):
        atomic_copy(REV / f, tmp / f)
        files.append(tmp / f)
    for g in re.findall(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}", tex):
        atomic_copy(REV / "figures" / g, tmp / g)
        files.append(tmp / g)
    return files


def git_bytes(rev: str, path: str) -> bytes:
    return subprocess.run(["git", "-C", str(REPO), "show", f"{rev}:{path}"], capture_output=True, check=True).stdout


def data_manifest() -> list[str]:
    """sha256 of the cited derived data as committed at TAG (bytes GitHub raw will serve; .gitattributes * -text)."""
    have_tag = subprocess.run(["git", "-C", str(REPO), "rev-parse", "-q", "--verify", f"refs/tags/{TAG}"],
                              capture_output=True).returncode == 0
    rev = TAG if have_tag else "HEAD"
    commit = subprocess.run(["git", "-C", str(REPO), "rev-parse", f"{rev}^{{commit}}"], capture_output=True,
                            text=True, check=True).stdout.strip()
    rows = []
    for path in DATA:
        b = git_bytes(rev, path)
        wd = TOPIC / path
        same = "=" if wd.exists() and wd.read_bytes() == b else ("(not in working folder)" if not wd.exists() else "DIFFERS")
        rows.append(f"| `{path}` | {len(b)} | `{hashlib.sha256(b).hexdigest()}` | {same} | |")
    names = subprocess.run(["git", "-C", str(REPO), "ls-tree", "--name-only", f"{rev}:results/awp_cache"],
                           capture_output=True, text=True, check=True).stdout.split()
    listing = "".join(f"{hashlib.sha256(git_bytes(rev, 'results/awp_cache/' + n)).hexdigest()}  {n}\n" for n in sorted(names))
    rows.append(f"| `results/awp_cache/*.npz` ({len(names)} files; sha256 of the sorted `sha256sum` listing) | — | "
                f"`{hashlib.sha256(listing.encode()).hexdigest()}` | | |")
    return ["", f"## C7(a) — derived data cited by the paper, as committed at `{rev}` (commit `{commit[:7]}`)", "",
            "Column *working folder*: `=` means the topic-folder copy is byte-identical. Column *after push*: after the "
            "author pushes the tag, fetch each file raw from GitHub at the tag "
            f"(`https://raw.githubusercontent.com/<owner>/<repo>/{TAG}/<path>`), hash it and fill in = / ≠. "
            "The URL in the paper currently points to the TPDS repository (TODO #6).", "",
            "| path | bytes | sha256 | working folder | after push (raw @ tag) |", "|---|---|---|---|---|"] + rows


def ascii_text(s: str) -> str:
    s = s.replace("---", " - ").replace("--", "-").replace("$", "")
    s = re.sub(r"\\(emph|textsc)\{([^}]*)\}", r"\2", s)
    rep = {r"\min\{\beta,c\}": "min{beta, c}", r"r^\star(c)": "r*(c)", r"r^\star(1)": "r*(1)", r"\beta": "beta",
           r"\{": "{", r"\}": "}"}
    for a, b in rep.items():
        s = s.replace(a, b)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return " ".join(s.split())


def main():
    PKG.mkdir(parents=True, exist_ok=True)
    for old in PKG.iterdir():  # regenerate numbered files + MANIFEST; keep hand-written guides/checklist
        if old.name[:2].isdigit() or old.name == "MANIFEST.md":
            old.unlink()
    atomic_copy(REV / "paper_v3.pdf", PKG / "01_manuscript_paper_v3.pdf")
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / "src"
        src.mkdir()
        files = flat_source(src)
        with zipfile.ZipFile(td / "02_source_flat.zip", "w", zipfile.ZIP_DEFLATED) as z:
            for f in files:
                z.write(f, f.name)
        atomic_copy(td / "02_source_flat.zip", PKG / "02_source_flat.zip")
        chk = td / "chk"
        chk.mkdir()
        with zipfile.ZipFile(PKG / "02_source_flat.zip") as z:
            z.extractall(chk)
        for i in range(3):
            subprocess.run(["pdflatex", "-interaction=nonstopmode", "paper_v3.tex"], cwd=chk,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        log = (chk / "paper_v3.log").read_text(encoding="utf-8", errors="ignore")
        errs = len(re.findall(r"(?m)^!", log))
        pages = re.search(r"Output written on paper_v3.pdf \((\d+) pages", log)
        build = f"flat source builds in an empty folder: {errs} errors, {pages.group(1) if pages else '?'} pages"
    atomic_copy(REV / "S1_theory_supplement.pdf", PKG / "03_S1_theory_supplement.pdf")
    atomic_copy(REV / "S2_empirical_supplement.pdf", PKG / "04_S2_empirical_supplement.pdf")
    atomic_copy(REV / "cover_letter.pdf", PKG / "05_cover_letter.pdf")
    import docx
    d = docx.Document()
    d.add_heading("Highlights", level=1)
    hl = [l.strip() for l in (REV / "highlights.txt").read_text(encoding="utf-8").splitlines() if l.strip()]
    for l in hl:
        d.add_paragraph(l, style="List Bullet")
    with atomic_open(PKG / "06_highlights.docx", mode="wb") as fh:
        d.save(fh)
    tpl = (REV / "paper_v3_template.tex").read_text(encoding="utf-8")
    title = re.search(r"\\title\{(.*?)\}\n", tpl).group(1)
    abstract = tpl[tpl.index("\\begin{abstract}") + 16: tpl.index("\\end{abstract}")]
    kw = re.search(r"\\begin\{keyword\}\n(.*?)\n\\end\{keyword\}", tpl, re.S).group(1).split("\\sep")
    txt = ["TITLE", ascii_text(title), "", "ABSTRACT (%d words)" % len(ascii_text(abstract).split()), ascii_text(abstract), "",
           "KEYWORDS", "; ".join(k.strip() for k in kw), "", "HIGHLIGHTS"] + ["- " + h for h in hl]
    atomic_write_text(PKG / "07_abstract_keywords_ASCII.txt", "\n".join(txt) + "\n", encoding="ascii")
    rows = [f"| {p.name} | {p.stat().st_size} | `{sha(p)}` |" for p in sorted(PKG.iterdir())]
    man = ["# MANIFEST — submission_package/EJOR (AwP v3)", "", f"Built by `code/v3_package.py`. {build}.", "",
           "| file | bytes | sha256 |", "|---|---|---|"] + rows + ["",
           "Not uploaded: MANIFEST.md, PORTAL_GUIDE_EJOR.md, CHECKLIST_03_v3.md (internal)."] + data_manifest()
    atomic_write_text(PKG / "MANIFEST.md", "\n".join(man) + "\n", encoding="utf-8")
    print(build)
    print("\n".join(rows))


if __name__ == "__main__":
    main()
